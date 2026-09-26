import io,json,zipfile
from unittest.mock import Mock,patch
from docx import Document as Word
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font
from django.test import TestCase,override_settings,Client
from django.core.files.uploadedfile import SimpleUploadedFile
from .tests import PortalTests
from .models import Document
from .office_engine import edit_office,inspect_office,OfficeError

def word_fixture():
    d=Word();d.add_heading('Office editing test',0)
    p=d.add_paragraph();p.add_run('Amount: USD ').bold=True;p.add_run('100').italic=True
    d.add_paragraph('Reference: ABC-4827');d.add_table(rows=1,cols=2).rows[0].cells[0].text='Keep table'
    b=io.BytesIO();d.save(b);return b.getvalue()
def excel_fixture():
    b=Workbook();s=b.active;s.title='Budget';s.append(['Item','Amount']);s.append(['Hosting',100]);s['B2'].font=Font(bold=True);s['B3']='=SUM(B2:B2)'
    out=io.BytesIO();b.save(out);return out.getvalue()

@override_settings(SECURE_SSL_REDIRECT=False)
class OfficeTests(TestCase):
    setUp=PortalTests.setUp
    def test_word_runs_and_unmodified_package(self):
        raw=word_fixture();out,changes=edit_office('source.docx',raw,{'operations':[{'type':'replace_text','find':'USD 100','replace':'USD 250'}]})
        d=Word(io.BytesIO(out));self.assertEqual(d.paragraphs[1].text,'Amount: USD 250');self.assertTrue(d.paragraphs[1].runs[0].bold);self.assertEqual(d.tables[0].cell(0,0).text,'Keep table')
        with zipfile.ZipFile(io.BytesIO(raw)) as a,zipfile.ZipFile(io.BytesIO(out)) as b:
            for name in a.namelist():
                if name!='word/document.xml':self.assertEqual(a.read(name),b.read(name))
        self.assertIn('USD 100',Word(io.BytesIO(raw)).paragraphs[1].text)
    def test_word_ambiguous_and_unsupported(self):
        for op in [{'type':'run_code'},{'type':'replace_text','find':'missing','replace':'x'},{'type':'set_paragraph','paragraph':999,'text':'x'}]:
            with self.assertRaises(OfficeError):edit_office('x.docx',word_fixture(),{'operations':[op]})
    def test_excel_preserves_formula_style_literal(self):
        out,_=edit_office('x.xlsx',excel_fixture(),{'operations':[{'type':'set_cell','sheet':'Budget','cell':'B2','value':250},{'type':'set_cell','sheet':'Budget','cell':'A4','value':'=not_a_formula','formula':False}]})
        b=load_workbook(io.BytesIO(out));self.assertEqual(b['Budget']['B2'].value,250);self.assertTrue(b['Budget']['B2'].font.bold);self.assertEqual(b['Budget']['B3'].value,'=SUM(B2:B2)');self.assertEqual(b['Budget']['A4'].data_type,'s')
        with self.assertRaises(OfficeError):edit_office('x.xlsx',excel_fixture(),{'operations':[{'type':'set_cell','sheet':'Budget','cell':'A4','value':'=WEBSERVICE("https://test")','formula':True}]})
    def test_api_ownership_csrf_and_original(self):
        for name,raw,operation in [('test.docx',word_fixture(),{'type':'replace_text','find':'USD 100','replace':'USD 250'}),('test.xlsx',excel_fixture(),{'type':'set_cell','sheet':'Budget','cell':'B2','value':250})]:
            r=self.client.post(f'/api/chats/{self.chat.id}/upload/',{'file':SimpleUploadedFile(name,raw)});pk=r.json()['id']
            self.assertEqual(self.client.get(f'/office/{pk}/').status_code,200)
            client=Client(enforce_csrf_checks=True);client.force_login(self.a);self.assertEqual(client.post(f'/api/office/{pk}/edit/',{},content_type='application/json').status_code,403)
            self.assertEqual(client.post(f'/api/office/{pk}/convert/',{},content_type='application/json').status_code,403)
            response=self.client.post(f'/api/office/{pk}/edit/',json.dumps({'operations':[operation]}),content_type='application/json');self.assertEqual(response.status_code,200,response.content)
            self.assertEqual(Document.objects.get(pk=pk).file.read(),raw)
            self.client.force_login(self.b)
            for path in [f'/office/{pk}/',f'/api/files/{response.json()["id"]}/']:self.assertEqual(self.client.get(path).status_code,404)
            for action in ['edit','convert']:self.assertEqual(self.client.post(f'/api/office/{pk}/{action}/',{},content_type='application/json').status_code,404)
            self.client.force_login(self.a)
    def test_ai_plan(self):
        r=self.client.post(f'/api/chats/{self.chat.id}/upload/',{'file':SimpleUploadedFile('test.docx',word_fixture())});pk=r.json()['id']
        mocked=Mock();mocked.json.return_value={'choices':[{'message':{'content':json.dumps({'operations':[{'type':'replace_text','find':'USD 100','replace':'USD 250'}]})}}]}
        with patch('portal.office_views.requests.post',return_value=mocked):
            r=self.client.post(f'/api/office/{pk}/edit/',json.dumps({'instruction':'Change USD 100 to USD 250'}),content_type='application/json')
        self.assertEqual(r.status_code,200,r.content)
