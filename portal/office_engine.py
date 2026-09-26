"""Bounded edits of Office files; preserve the original package for Word edits."""
import io,json,math,re,zipfile
from pathlib import Path
from lxml import etree as ET
from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
W='{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
class OfficeError(ValueError): pass

def package(raw):
    try:
        z=zipfile.ZipFile(io.BytesIO(raw)); items=z.infolist()
        if len(items)>3000 or sum(i.file_size for i in items)>80*1024*1024: raise OfficeError('The expanded file is too large.')
        names=z.namelist()
        if len(set(names))!=len(names): raise OfficeError('Invalid Office package.')
        for name in names:
            low=name.lower()
            if any(x in low for x in ('vbaproject','activex/','embeddings/','externallinks/','connections.xml','_xmlsignatures/')): raise OfficeError('Macros, embedded objects, external data connections and signed files are not supported.')
            if low.endswith(('.xml','.rels')):
                xml=z.read(name)
                if b'<!DOCTYPE' in xml.upper() or b'<!ENTITY' in xml.upper():raise OfficeError('Unsupported XML entities.')
                root=ET.fromstring(xml,ET.XMLParser(resolve_entities=False,no_network=True))
                for el in root.iter():
                    if el.get('TargetMode')=='External' and not el.get('Type','').endswith('/hyperlink'):raise OfficeError('Remove linked external resources before editing or converting.')
                    if el.tag in (W+'instrText',W+'fldSimple'):
                        value=(el.text or '')+' '+el.get(W+'instr','')
                        if re.search(r'\b(DDE|DDEAUTO|INCLUDETEXT|INCLUDEPICTURE|LINK|DATABASE)\b',value,re.I):raise OfficeError('Linked document fields are not supported.')
        return z
    except OfficeError:raise
    except Exception as exc:raise OfficeError('This Office file is invalid or unsupported.') from exc

def word_parts(z):
    result=[]
    for name in z.namelist():
        if re.fullmatch(r'word/(document|header\d+|footer\d+|footnotes|endnotes)\.xml',name):
            root=ET.fromstring(z.read(name),ET.XMLParser(resolve_entities=False,no_network=True))
            if root.find('.//'+W+'ins') is not None or root.find('.//'+W+'del') is not None:raise OfficeError('Accept or reject tracked changes in Word before editing.')
            result.append((name,root))
    return result

def paragraphs(parts):
    return [(root,p) for _,root in parts for p in root.iter(W+'p')]
def ptext(p):return ''.join(n.text or '' for n in p.iter(W+'t'))

def inspect_office(name,raw):
    ext=Path(name).suffix.lower();z=package(raw)
    if ext=='.docx':
        rows=[];size=0
        for index,(_,p) in enumerate(paragraphs(word_parts(z)),1):
            text=ptext(p)
            if text: rows.append({'paragraph':index,'text':text[:2500]});size+=len(text)
            if len(rows)>=250 or size>24000:break
        return {'kind':'word','paragraphs':rows,'note':'Paragraph numbers refer to the original file. Preview may be abbreviated.'}
    if ext!='.xlsx':raise OfficeError('Use DOCX or XLSX files.')
    book=load_workbook(io.BytesIO(raw),read_only=False,data_only=False,keep_links=False)
    sheets=[];total=0
    for sheet in book:
        cells=[]
        for row in sheet.iter_rows(max_row=min(sheet.max_row,500),max_col=min(sheet.max_column,40)):
            for cell in row:
                if cell.value is not None:
                    cells.append({'cell':cell.coordinate,'value':str(cell.value)[:500]});total+=1
                    if total>=600:break
            if total>=600:break
        sheets.append({'sheet':sheet.title,'rows':sheet.max_row,'columns':sheet.max_column,'cells':cells})
        if total>=600:break
    book.close()
    return {'kind':'excel','sheets':sheets,'note':'Preview includes at most 600 non-empty cells, within the first 500 rows and 40 columns. Formulas are shown, not evaluated.'}

def replace_span(p,start,end,value):
    nodes=list(p.iter(W+'t'));pos=0;inserted=False
    for n in nodes:
        text=n.text or '';a,b=pos,pos+len(text);pos=b
        if b<=start or a>=end:continue
        prefix=text[:max(0,start-a)];suffix=text[max(0,end-a):] if end<b else ''
        n.text=prefix+(value if not inserted else '')+suffix;n.set('{http://www.w3.org/XML/1998/namespace}space','preserve');inserted=True

def edit_office(name,raw,plan):
    ops=plan.get('operations') if isinstance(plan,dict) else None
    if isinstance(plan,dict) and plan.get('error'):raise OfficeError(str(plan['error'])[:500])
    if not isinstance(ops,list) or not 1<=len(ops)<=50:raise OfficeError('Specify between 1 and 50 edits.')
    z=package(raw);ext=Path(name).suffix.lower();changes=[]
    if ext=='.docx':
        parts=word_parts(z);paras=paragraphs(parts)
        for op in ops:
            if not isinstance(op,dict):raise OfficeError('Invalid operation.')
            kind=op.get('type');value=op.get('replace',op.get('text',''))
            if not isinstance(value,str) or len(value)>10000 or any(ord(c)<32 and c not in '\n\t' for c in value):raise OfficeError('Invalid replacement text.')
            if '\n' in value or '\t' in value:raise OfficeError('Use one paragraph per edit; line breaks and tabs are not supported.')
            if kind=='replace_text':
                find=op.get('find')
                if not isinstance(find,str) or not find or len(find)>10000:raise OfficeError('Enter the exact text to replace.')
                targets=paras
                if 'paragraph' in op:
                    index=op['paragraph']
                    if type(index)!=int or not 1<=index<=len(paras):raise OfficeError('Invalid paragraph number.')
                    targets=[paras[index-1]]
                matches=[(p,m.start(),m.end()) for _,p in targets for m in re.finditer(re.escape(find),ptext(p))]
                if not matches:raise OfficeError('The requested text was not found in a supported paragraph.')
                if len(matches)>1 and op.get('all') is not True:raise OfficeError('The text occurs more than once. Specify a paragraph or choose Replace all.')
                for p,a,b in reversed(matches):replace_span(p,a,b,value)
                changes.append(f'Replaced {len(matches)} occurrence(s).')
            elif kind=='set_paragraph':
                index=op.get('paragraph')
                if type(index)!=int or not 1<=index<=len(paras):raise OfficeError('Invalid paragraph number.')
                p=paras[index-1][1];old=ptext(p)
                if not old:raise OfficeError('Select a paragraph containing text.')
                replace_span(p,0,len(old),value);changes.append(f'Updated paragraph {index}.')
            else:raise OfficeError('Supported Word edits: replace_text and set_paragraph.')
        changed={name:ET.tostring(root,xml_declaration=True,encoding='UTF-8',standalone=True) for name,root in parts}
        output=io.BytesIO()
        with zipfile.ZipFile(output,'w',zipfile.ZIP_DEFLATED) as out:
            for item in z.infolist():out.writestr(item,changed.get(item.filename,z.read(item.filename)))
        return output.getvalue(),changes
    if ext!='.xlsx':raise OfficeError('Use DOCX or XLSX files.')
    book=load_workbook(io.BytesIO(raw),data_only=False,keep_links=False)
    try:
        for op in ops:
            if not isinstance(op,dict) or op.get('type')!='set_cell':raise OfficeError('Supported Excel edit: set_cell.')
            sheet=op.get('sheet');address=op.get('cell','');value=op.get('value');formula=op.get('formula',False)
            if sheet not in book.sheetnames or not isinstance(address,str) or not re.fullmatch(r'[A-Z]{1,3}[1-9][0-9]{0,5}',address):raise OfficeError('Specify an existing sheet and valid cell address.')
            cell=book[sheet][address]
            if cell.row>100000 or cell.column>1000 or isinstance(cell,MergedCell):raise OfficeError('Cell is outside supported bounds or inside a merged area.')
            if not isinstance(value,(str,int,float,bool,type(None))) or (isinstance(value,float) and not math.isfinite(value)):raise OfficeError('Invalid cell value.')
            if isinstance(value,str) and len(value)>10000:raise OfficeError('Cell text is too long.')
            if formula:
                if not isinstance(value,str) or not value.startswith('=') or len(value)>1000:raise OfficeError('Formulas must begin with = and contain at most 1,000 characters.')
                if re.search(r'[\[\]\\|]|https?:|file:|\b(WEBSERVICE|HYPERLINK|RTD|DDE|CALL|REGISTER|IMAGE)\s*\(',value,re.I):raise OfficeError('External or executable formulas are not supported.')
            cell.value=value
            if isinstance(value,str) and not formula:cell.data_type='s'
            changes.append(f'Updated {sheet}!{address}.')
        from openpyxl.workbook.properties import CalcProperties
        book.calculation=CalcProperties(calcId=0,fullCalcOnLoad=True,forceFullCalc=True)
        output=io.BytesIO();book.save(output);return output.getvalue(),changes
    finally:book.close()
