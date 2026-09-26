import io, tempfile, os
from unittest.mock import patch, Mock
from django.test import TestCase, Client, override_settings
from django.contrib.auth import get_user_model
from django.core.cache import cache
from django.core.files.uploadedfile import SimpleUploadedFile
from .models import Conversation, Document, Message
from .documents import extract

@override_settings(SECURE_SSL_REDIRECT=False)
class PortalTests(TestCase):
    def setUp(self):
        cache.clear()
        self.temp=tempfile.TemporaryDirectory()
        self.media=override_settings(MEDIA_ROOT=self.temp.name); self.media.enable()
        self.addCleanup(self.temp.cleanup); self.addCleanup(self.media.disable)
        self.a=get_user_model().objects.create_user('alice',password='testing-password-529')
        self.b=get_user_model().objects.create_user('bob',password='testing-password-783')
        self.chat=Conversation.objects.create(owner=self.a)
        self.client.force_login(self.a)
    def test_user_isolation(self):
        doc=Document.objects.create(owner=self.a,conversation=self.chat,name='private.txt',size=4,text='data',file=SimpleUploadedFile('p.txt',b'data'))
        self.client.force_login(self.b)
        for url in [f'/api/chats/{self.chat.id}/',f'/api/files/{doc.id}/']:
            self.assertEqual(self.client.get(url).status_code,404)
        self.assertEqual(self.client.post(f'/api/chats/{self.chat.id}/send/',data='{"message":"hello"}',content_type='application/json').status_code,404)
        self.assertEqual(self.client.post(f'/api/chats/{self.chat.id}/upload/',{'file':SimpleUploadedFile('x.txt',b'secret')}).status_code,404)
    def test_regular_user_cannot_admin(self):
        self.assertEqual(self.client.get('/admin/auth/user/').status_code,302)
        self.assertEqual(self.client.post('/setup/',{}).status_code,404)
    def test_disabled_account_loses_session(self):
        self.a.is_active=False; self.a.save()
        self.assertEqual(self.client.get('/').status_code,302)
    def test_csrf_required_for_upload_and_message(self):
        client=Client(enforce_csrf_checks=True);client.force_login(self.a)
        for path in [f'/api/chats/{self.chat.id}/send/',f'/api/chats/{self.chat.id}/upload/','/api/chats/']:
            self.assertEqual(client.post(path,{}).status_code,403)
    def test_stale_login_form_recovers_without_authenticating(self):
        client=Client(enforce_csrf_checks=True)
        page=client.get('/login/')
        self.assertIn('no-store', page['Cache-Control'])
        stale=client.cookies['csrftoken'].value
        from django.middleware.csrf import _get_new_csrf_string
        client.cookies['csrftoken']=_get_new_csrf_string()
        response=client.post('/login/',{'username':'alice','password':'testing-password-529','csrfmiddlewaretoken':stale})
        self.assertRedirects(response,'/login/?expired=1')
        self.assertNotIn('_auth_user_id',client.session)
        page=client.get(response.url)
        self.assertContains(page,'This sign-in page has expired')
        fresh=client.cookies['csrftoken'].value
        response=client.post('/login/',{'username':'alice','password':'testing-password-529','csrfmiddlewaretoken':fresh})
        self.assertEqual(response.status_code,302)
        self.assertIn('_auth_user_id',client.session)
    def test_login_throttled_and_username_works(self):
        self.client.logout()
        self.assertEqual(self.client.post('/login/',{'username':'alice','password':'testing-password-529'}).status_code,302)
        self.client.logout()
        for _ in range(10): self.client.post('/login/',{'username':'alice','password':'wrong'})
        self.assertEqual(self.client.post('/login/',{'username':'alice','password':'wrong'}).status_code,429)
    def test_streaming_answer_saves_only_completed_exchange(self):
        upstream=Mock()
        upstream.iter_lines.return_value=iter([b'data: {"choices":[{"delta":{"content":"hello"},"finish_reason":null}]}', b'data: {"choices":[{"delta":{},"finish_reason":"stop"}]}',b'data: [DONE]'])
        with patch('portal.streaming.requests.post',return_value=upstream):
            response=self.client.post(f'/api/chats/{self.chat.id}/send/',data='{"message":"Hi"}',content_type='application/json',HTTP_ACCEPT='application/x-ndjson')
            self.assertTrue(response.streaming)
            it=iter(response.streaming_content)
            self.assertIn(b'"start"',next(it))
            self.assertIn(b'hello',next(it))
            self.assertEqual(self.chat.messages.count(),0)
            self.assertIn(b'"done"',b''.join(it))
        self.assertEqual(self.chat.messages.count(),2)
        upstream.close.assert_called_once()
        from .views import _generation_lock
        self.assertFalse(_generation_lock.locked())
    def test_interrupted_stream_releases_lock_without_saving(self):
        upstream=Mock()
        upstream.iter_lines.return_value=iter([b'data: {"choices":[{"delta":{"content":"partial"}}]}'])
        with patch('portal.streaming.requests.post',return_value=upstream):
            response=self.client.post(f'/api/chats/{self.chat.id}/send/',data='{"message":"Hi"}',content_type='application/json',HTTP_ACCEPT='application/x-ndjson')
            self.assertIn(b'"error"',b''.join(response.streaming_content))
        self.assertEqual(self.chat.messages.count(),0)
        from .views import _generation_lock
        self.assertFalse(_generation_lock.locked())
    def test_report_can_stream_beyond_old_100_second_deadline(self):
        upstream=Mock()
        upstream.iter_lines.return_value=iter([
            b'data: {"choices":[{"delta":{"content":"long report complete"},"finish_reason":"stop"}]}'])
        with patch('portal.streaming.requests.post',return_value=upstream), patch('portal.streaming.time.monotonic',side_effect=[0,150]):
            response=self.client.post(f'/api/chats/{self.chat.id}/send/',data='{"message":"Report"}',content_type='application/json',HTTP_ACCEPT='application/x-ndjson')
            self.assertIn(b'"done"',b''.join(response.streaming_content))
        self.assertEqual(self.chat.messages.count(),2)

    def test_length_limited_report_is_saved_as_incomplete(self):
        import json
        upstream=Mock()
        upstream.iter_lines.return_value=iter([
            b'data: {"choices":[{"delta":{"content":"report cut here"},"finish_reason":"length"}]}',
            b'data: [DONE]'])
        with patch('portal.streaming.requests.post',return_value=upstream) as post:
            response=self.client.post(f'/api/chats/{self.chat.id}/send/',data='{"message":"Analyze report"}',content_type='application/json',HTTP_ACCEPT='application/x-ndjson')
            events=[json.loads(line) for line in b''.join(response.streaming_content).splitlines()]
            self.assertTrue(events[-1]['incomplete'])
            self.assertEqual(post.call_args.kwargs['json']['max_tokens'],4096)
        saved=self.chat.messages.get(role='assistant')
        self.assertTrue(saved.metadata['incomplete'])
        restored=self.client.get(f'/api/chats/{self.chat.id}/').json()
        self.assertTrue(restored['messages'][-1]['incomplete'])

    def test_long_answer_tail_is_available_for_continuation(self):
        Message.objects.create(conversation=self.chat,role='assistant',content='a'*9000+'LAST ROW')
        upstream=Mock()
        upstream.iter_lines.return_value=iter([
            b'data: {"choices":[{"delta":{"content":"remainder"},"finish_reason":"stop"}]}'])
        with patch('portal.streaming.requests.post',return_value=upstream) as post:
            response=self.client.post(f'/api/chats/{self.chat.id}/send/',data='{"message":"Continue the previous response"}',content_type='application/json',HTTP_ACCEPT='application/x-ndjson')
            b''.join(response.streaming_content)
            messages=post.call_args.kwargs['json']['messages']
            self.assertTrue(messages[-2]['content'].endswith('LAST ROW'))

    def test_done_marker_without_finish_reason_is_not_success(self):
        upstream=Mock()
        upstream.iter_lines.return_value=iter([
            b'data: {"choices":[{"delta":{"content":"partial"}}]}',b'data: [DONE]'])
        with patch('portal.streaming.requests.post',return_value=upstream):
            response=self.client.post(f'/api/chats/{self.chat.id}/send/',data='{"message":"Hi"}',content_type='application/json',HTTP_ACCEPT='application/x-ndjson')
            self.assertIn(b'"error"',b''.join(response.streaming_content))
        self.assertEqual(self.chat.messages.count(),0)

    def test_image_upload_preview_and_model_input(self):
        from PIL import Image
        buf=io.BytesIO(); Image.new('RGB',(2000,1000),'red').save(buf,format='PNG')
        response=self.client.post(f'/api/chats/{self.chat.id}/upload/',{'file':SimpleUploadedFile('sample.png',buf.getvalue(),content_type='image/png')})
        self.assertEqual(response.status_code,200)
        self.assertTrue(response.json()['image'])
        doc=Document.objects.get(id=response.json()['id'])
        with doc.file.open('rb') as f:
            picture=Image.open(f); self.assertEqual(picture.size,(1280,640)); self.assertEqual(picture.format,'JPEG')
        preview=self.client.get(f'/api/files/{doc.id}/?preview=1'); self.assertEqual(preview['Content-Type'],'image/jpeg'); preview.close()
        reply=Mock();reply.json.return_value={'choices':[{'message':{'content':'红色'}}]}
        with patch('portal.views.requests.post',return_value=reply) as call:
            response=self.client.post(f'/api/chats/{self.chat.id}/send/',data='{"message":"what color?"}',content_type='application/json')
        self.assertEqual(response.status_code,200)
        content=call.call_args.kwargs['json']['messages'][-1]['content']
        self.assertTrue(content[1]['image_url']['url'].startswith('data:image/jpeg;base64,'))
        self.client.force_login(self.b)
        self.assertEqual(self.client.get(f'/api/files/{doc.id}/?preview=1').status_code,404)
    def test_invalid_image_and_image_limit(self):
        from PIL import Image
        bad=self.client.post(f'/api/chats/{self.chat.id}/upload/',{'file':SimpleUploadedFile('fake.png',b'not an image')})
        self.assertEqual(bad.status_code,400)
        buf=io.BytesIO();Image.new('RGB',(32,32),'blue').save(buf,format='PNG')
        for i in range(3):
            response=self.client.post(f'/api/chats/{self.chat.id}/upload/',{'file':SimpleUploadedFile('image.png',buf.getvalue())})
            self.assertEqual(response.status_code,200 if i<2 else 400)
    def test_upload_and_model_context(self):
        r=self.client.post(f'/api/chats/{self.chat.id}/upload/',{'file':SimpleUploadedFile('plan.txt','项目代号是青竹，交付日期为10月8日。'.encode())})
        self.assertEqual(r.status_code,200)
        response=Mock(); response.json.return_value={'choices':[{'message':{'content':'项目代号是青竹。'}}]}
        with patch('portal.views.requests.post',return_value=response) as call:
            r=self.client.post(f'/api/chats/{self.chat.id}/send/',data='{"message":"项目代号是什么？"}',content_type='application/json')
        self.assertEqual(r.status_code,200)
        self.assertIn('青竹',call.call_args.kwargs['json']['messages'][0]['content'])
        self.assertEqual(Message.objects.filter(conversation=self.chat).count(),2)
    def test_failed_model_does_not_save_incomplete_exchange(self):
        with patch('portal.views.requests.post',side_effect=TimeoutError):
            r=self.client.post(f'/api/chats/{self.chat.id}/send/',data='{"message":"你好"}',content_type='application/json')
        self.assertEqual(r.status_code,503);self.assertEqual(self.chat.messages.count(),0)
    def test_reject_executable_upload(self):
        r=self.client.post(f'/api/chats/{self.chat.id}/upload/',{'file':SimpleUploadedFile('evil.html',b'<script>alert(1)</script>')})
        self.assertEqual(r.status_code,400)
    @override_settings(ENABLE_SETUP=True)
    def test_setup_only_once(self):
        data={'username':'manager','password1':'Strong-test-pass-9641','password2':'Strong-test-pass-9641'}
        self.assertEqual(self.client.post('/setup/',data).status_code,200)
        self.assertTrue(get_user_model().objects.get(username='manager').is_superuser)
        self.assertEqual(self.client.get('/setup/').status_code,404)
    def test_document_parsers(self):
        from docx import Document as WordDocument
        from openpyxl import Workbook
        doc=WordDocument();doc.add_paragraph('合同总额：25000元');buf=io.BytesIO();doc.save(buf)
        self.assertIn('25000',extract('example.docx',buf.getvalue())[0])
        wb=Workbook();wb.active.append(['产品','数量']);wb.active.append(['测试',42]);buf=io.BytesIO();wb.save(buf)
        self.assertIn('42',extract('example.xlsx',buf.getvalue())[0])
        self.assertIn('测试',extract('example.csv','产品,数量\n测试,42'.encode())[0])
    def test_runpod_proxy_preserves_host_boundary(self):
        from . import wsgi
        with patch.dict(os.environ,{'RUNPOD_POD_ID':'testpod123','PORTAL_SETUP':'0'}), patch.object(wsgi,'_django_application',return_value=[]):
            internal={'HTTP_HOST':'100.65.18.208:60121'}
            wsgi.application(internal,lambda *args:None)
            self.assertEqual(internal['HTTP_HOST'],'testpod123-8090.proxy.runpod.net')
            external={'HTTP_HOST':'untrusted.example'}
            wsgi.application(external,lambda *args:None)
            self.assertEqual(external['HTTP_HOST'],'untrusted.example')
